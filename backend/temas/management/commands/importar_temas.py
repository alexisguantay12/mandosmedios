import hashlib
import json
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from temas.models import (
    ActividadTema,
    Area,
    Comentario,
    Estado,
    Responsable,
    Tema,
    TipoTema,
)


# ============================================================
# CONFIGURACIÓN DEL FORMULARIO HISTÓRICO
# ============================================================

HOJA_DEFAULT = "Respuestas de formulario 1"

FUENTE = "formulario-mandos-medios-historico"


# ============================================================
# NOMBRES DE COLUMNAS
# ============================================================

COL_FECHA = "Marca temporal"

COL_TIPO = "Tipo de planteo"

COL_INFORMATIVO = (
    "Cuenta brevemente qué querés informar. "
    "Podés usar este espacio para comunicar gestiones en curso, "
    "novedades del sector, buenas prácticas o agradecimientos. "
    "No se generará seguimiento específico."
)

COL_REMITENTE = "Nombre del remitente"

COL_AREA = "Área o sector"

COL_PROBLEMA = (
    "Problema planteado "
    "Describí brevemente la situación o necesidad detectada."
)

COL_IMPACTO = "Impacto Describe de qué manera afecta"

COL_SOLUCION = (
    "Soluciones propuestas "
    "¿Qué acciones o ideas sugerís para abordarlo o resolverlo?"
)

COL_ESTADO_ACTUAL = (
    "Estado actual "
    "¿Qué se ha hecho hasta ahora? "
    "¿Hay gestiones iniciadas o intentos previos de resolución?"
)

COL_REMITENTE_2 = "Nombre del remitente 2"

COL_COMENTARIOS_REUNION = "Comentarios en reunión"

COL_ESTADO_TEMA = "Estado del tema"

COL_FECHA_LIMITE = "Fecha límite de acción"


# ============================================================
# RESPONSABLES DEL FORMULARIO HISTÓRICO
#
# En el Excel están como "Apellido Nombre".
# En el sistema están como "Nombre Apellido".
# ============================================================

COLUMNAS_RESPONSABLES = [
    "Burgos Blanca",
    "Burgos Nancy",
    "Cabrera Alejandro",
    "Ferrufino Lila",
    "Avellaneda Emanuel",
    "Guantay Alexis",
    "Gudiño Ana",
    "Madariaga Alicia",
    "Mercado Evangelina",
    "Paez Debora",
    "Ram Natalia",
    "Ulloa Paola",
    "Zuviria Victoria",
]


# ============================================================
# ALIASES DE ÁREAS
#
# Izquierda: valor histórico normalizado.
# Derecha: nombre actual del maestro Area.
# ============================================================

ALIASES_AREAS = {
    # Administración
    "administracion": "Administracion",

    # Contable
    "contable": "Contable",
    "contables": "Contable",
    "contabilidad": "Contable",

    # Enfermería
    "enfermeria": "Enfermeria",

    # Gestión del paciente
    "gestion del paciente": "Gestion del Paciente",

    # Intendencia
    "intendencia": "Intendencia",

    # Mantenimiento
    "mantenimiento": "Mantenimiento",

    # Portería / Seguridad
    "porteria": "Porteria",
    "seguridad": "Porteria",

    # Quirófano
    "quirofano": "Quirofano",

    # Recursos Humanos
    "rrhh": "Recursos Humanos",
    "recursos humanos": "Recursos Humanos",

    # Sistemas
    "sistemas": "Sistemas",
}


# ============================================================
# ÁREAS HISTÓRICAS QUE NO EXISTEN EN TU MAESTRO ACTUAL
#
# Estas áreas son realmente distintas. No corresponde inventar
# que "Farmacia" sea Sistemas, Gestión del Paciente, etc.
#
# El importador las crea solamente si aparecen.
# ============================================================

AREAS_HISTORICAS_NUEVAS = {
    "farmacia": {
        "nombre": "Farmacia",
        "color": "#BD577D",
    },
    "nutricion": {
        "nombre": "Nutricion",
        "color": "#8A9A5B",
    },
    "direccion medica": {
        "nombre": "Direccion Medica",
        "color": "#496D8A",
    },
    "otros": {
        "nombre": "Otros",
        "color": "#7A8491",
    },
}


# ============================================================
# RESPONSABLES
#
# Sólo se usan responsables ya existentes.
#
# Natalia Ram aparece históricamente en el Excel, pero si no
# existe actualmente como Responsable, NO se crea automáticamente.
# Queda preservada en datos_migracion.
# ============================================================

ALIASES_RESPONSABLES = {
    "burgos blanca": "Blanca Burgos",
    "burgos nancy": "Nancy Burgos",
    "cabrera alejandro": "Alejandro Cabrera",
    "ferrufino lila": "Lila Ferrufino",
    "avellaneda emanuel": "Emanuel Avellaneda",
    "guantay alexis": "Alexis Guantay",
    "gudino ana": "Ana Gudiño",
    "madariaga alicia": "Alicia Madariaga",
    "mercado evangelina": "Evangelina Mercado",
    "paez debora": "Debora Paez",
    "ram natalia": "Natalia Ram",
    "ulloa paola": "Paola Ulloa",
    "zuviria victoria": "Victoria Zuviria",
}


# ============================================================
# UTILIDADES
# ============================================================

def texto(value):
    if value is None:
        return ""

    return str(value).replace("\r\n", "\n").strip()


def normalizar(value):
    """
    Normalización para comparar:
    - mayúsculas/minúsculas
    - acentos
    - espacios
    - saltos de línea
    """
    value = texto(value)

    value = unicodedata.normalize(
        "NFKD",
        value,
    )

    value = "".join(
        c
        for c in value
        if not unicodedata.combining(c)
    )

    return " ".join(
        value.casefold().split()
    )


def serializable(value):
    if isinstance(
        value,
        (
            datetime,
            date,
            time,
        ),
    ):
        return value.isoformat()

    return value


def celda_es_true(value):
    """
    Google Forms / Excel puede devolver boolean,
    número o texto.
    """

    if value is True:
        return True

    if isinstance(value, (int, float)):
        return value == 1

    value = normalizar(value)

    return value in {
        "true",
        "verdadero",
        "si",
        "sí",
        "x",
        "1",
    }


def interpretar_fecha(value, epoch):
    if value is None:
        return None

    if texto(value) == "":
        return None

    if isinstance(value, (int, float)):
        value = from_excel(
            value,
            epoch,
        )

    if isinstance(value, str):
        formatos = [
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d/%m/%Y",
            "%d-%m-%Y",
            "%Y-%m-%d",
        ]

        parsed = None

        try:
            parsed = datetime.fromisoformat(
                value.replace(
                    "Z",
                    "+00:00",
                )
            )
        except ValueError:
            pass

        if parsed is None:
            for formato in formatos:
                try:
                    parsed = datetime.strptime(
                        value.strip(),
                        formato,
                    )
                    break
                except ValueError:
                    continue

        if parsed is None:
            raise ValueError(
                f"Fecha no interpretable: {value!r}"
            )

        value = parsed

    if isinstance(value, date) and not isinstance(
        value,
        datetime,
    ):
        value = datetime.combine(
            value,
            time.min,
        )

    if not isinstance(value, datetime):
        raise ValueError(
            f"Fecha no válida: {value!r}"
        )

    if timezone.is_naive(value):
        value = timezone.make_aware(
            value
        )

    return value


def intentar_fecha(value, epoch):
    """
    La columna 'Fecha límite de acción' no siempre contiene
    una fecha. A veces contiene texto.

    Devuelve:
        (fecha, texto_no_fecha)
    """

    if not texto(value):
        return None, ""

    try:
        fecha = interpretar_fecha(
            value,
            epoch,
        )

        return fecha.date(), ""

    except (ValueError, TypeError):
        return None, texto(value)


def generar_titulo(descripcion):
    """
    Genera un título legible sin perder la descripción completa.
    """

    limpio = " ".join(
        texto(descripcion).split()
    )

    if not limpio:
        return "Tema importado"

    limite = 110

    if len(limpio) <= limite:
        return limpio

    titulo = limpio[:limite]

    if " " in titulo:
        titulo = titulo.rsplit(
            " ",
            1,
        )[0]

    return titulo + "…"


# ============================================================
# COMMAND
# ============================================================

class Command(BaseCommand):
    help = (
        "Importa el formulario histórico de Mandos Medios "
        "al sistema actual."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "archivo",
            help="Ruta del archivo XLSX.",
        )

        parser.add_argument(
            "--usuario",
            required=True,
            help=(
                "Username Django que figurará como usuario "
                "importador."
            ),
        )

        parser.add_argument(
            "--sheet",
            default=HOJA_DEFAULT,
            help=(
                "Hoja a importar. "
                f"Default: {HOJA_DEFAULT}"
            ),
        )

        parser.add_argument(
            "--dry-run",
            action="store_true",
            help=(
                "Procesa todo pero realiza rollback al finalizar."
            ),
        )


    # ========================================================
    # HANDLE
    # ========================================================

    def handle(self, *args, **options):
        path = Path(
            options["archivo"]
        )

        if not path.exists():
            raise CommandError(
                f"No existe el archivo: {path}"
            )

        if path.suffix.lower() != ".xlsx":
            raise CommandError(
                "El archivo debe ser .xlsx"
            )

        try:
            usuario = (
                get_user_model()
                .objects
                .get(
                    username=options["usuario"],
                    is_active=True,
                )
            )

        except get_user_model().DoesNotExist as exc:
            raise CommandError(
                "El usuario importador no existe o está inactivo."
            ) from exc


        try:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise CommandError(
                f"No se pudo leer el Excel: {exc}"
            ) from exc


        sheet_name = options["sheet"]

        if sheet_name not in workbook.sheetnames:
            disponibles = ", ".join(
                workbook.sheetnames
            )

            workbook.close()

            raise CommandError(
                f"La hoja {sheet_name!r} no existe. "
                f"Disponibles: {disponibles}"
            )


        sheet = workbook[
            sheet_name
        ]


        resumen = {
            "procesadas": 0,
            "creadas": 0,
            "omitidas": 0,
            "errores": 0,
            "comentarios": 0,
            "areas_creadas": [],
            "responsables_no_encontrados": set(),
        }


        try:
            with transaction.atomic():

                self.importar_hoja(
                    workbook,
                    sheet,
                    path,
                    usuario,
                    resumen,
                )

                if options["dry_run"]:
                    transaction.set_rollback(
                        True
                    )

        finally:
            workbook.close()


        self.stdout.write("")
        self.stdout.write(
            "=" * 60
        )

        self.stdout.write(
            "RESUMEN DE IMPORTACIÓN"
        )

        self.stdout.write(
            "=" * 60
        )

        self.stdout.write(
            f"Procesadas: {resumen['procesadas']}"
        )

        self.stdout.write(
            f"Creadas: {resumen['creadas']}"
        )

        self.stdout.write(
            f"Omitidas por ya existir: {resumen['omitidas']}"
        )

        self.stdout.write(
            f"Errores: {resumen['errores']}"
        )

        self.stdout.write(
            f"Comentarios históricos: {resumen['comentarios']}"
        )


        if resumen["areas_creadas"]:
            self.stdout.write("")
            self.stdout.write(
                "Áreas históricas creadas:"
            )

            for nombre in sorted(
                set(
                    resumen[
                        "areas_creadas"
                    ]
                )
            ):
                self.stdout.write(
                    f"  - {nombre}"
                )


        if resumen[
            "responsables_no_encontrados"
        ]:
            self.stdout.write("")

            self.stdout.write(
                self.style.WARNING(
                    "Responsables históricos sin maestro actual:"
                )
            )

            for nombre in sorted(
                resumen[
                    "responsables_no_encontrados"
                ]
            ):
                self.stdout.write(
                    f"  - {nombre}"
                )


        if options["dry_run"]:
            self.stdout.write("")

            self.stdout.write(
                self.style.WARNING(
                    "DRY RUN finalizado: "
                    "NO se guardó ningún cambio."
                )
            )

        else:
            self.stdout.write("")

            self.stdout.write(
                self.style.SUCCESS(
                    "Importación finalizada correctamente."
                )
            )


    # ========================================================
    # IMPORTAR HOJA
    # ========================================================

    def importar_hoja(
        self,
        workbook,
        sheet,
        path,
        usuario,
        resumen,
    ):
        iterator = sheet.iter_rows(
            values_only=True
        )

        raw_headers = next(
            iterator,
            None,
        )

        if not raw_headers:
            raise CommandError(
                "La hoja no contiene encabezados."
            )


        headers = [
            texto(h)
            for h in raw_headers
        ]


        indexes = {}

        for index, header in enumerate(
            headers
        ):
            if not header:
                continue

            key = normalizar(
                header
            )

            indexes[key] = index


        # Buscar encabezados por fragmentos para no depender
        # de saltos de línea / espacios exactos.

        def buscar_columna(
            fragmento,
            obligatoria=True,
        ):
            fragmento = normalizar(
                fragmento
            )

            matches = [
                index
                for key, index
                in indexes.items()
                if fragmento in key
            ]

            if matches:
                return matches[0]

            if obligatoria:
                raise CommandError(
                    f"No se encontró columna que contenga: "
                    f"{fragmento!r}"
                )

            return None


        columnas = {
            "fecha":
                buscar_columna(
                    "marca temporal"
                ),

            "tipo":
                buscar_columna(
                    "tipo de planteo"
                ),

            "informativo":
                buscar_columna(
                    "cuenta brevemente"
                ),

            "remitente":
                buscar_columna(
                    "nombre del remitente"
                ),

            "area":
                buscar_columna(
                    "area o sector"
                ),

            "problema":
                buscar_columna(
                    "problema planteado"
                ),

            "impacto":
                buscar_columna(
                    "impacto"
                ),

            "solucion":
                buscar_columna(
                    "soluciones propuestas"
                ),

            "estado_actual":
                buscar_columna(
                    "estado actual"
                ),

            "comentario_reunion":
                buscar_columna(
                    "comentarios en reunion"
                ),

            "estado":
                buscar_columna(
                    "estado del tema"
                ),

            "fecha_limite":
                buscar_columna(
                    "fecha limite de accion"
                ),
        }


        columnas_responsables = {}

        for nombre in COLUMNAS_RESPONSABLES:
            key = normalizar(
                nombre
            )

            if key in indexes:
                columnas_responsables[
                    nombre
                ] = indexes[key]


        for rownum, row in enumerate(
            iterator,
            start=2,
        ):
            if not any(
                texto(value)
                for value in row
            ):
                continue

            resumen[
                "procesadas"
            ] += 1

            try:
                with transaction.atomic():
                    creada = self.importar_fila(
                        workbook=workbook,
                        sheet=sheet,
                        path=path,
                        row=row,
                        rownum=rownum,
                        headers=headers,
                        columnas=columnas,
                        columnas_responsables=columnas_responsables,
                        usuario=usuario,
                        resumen=resumen,
                    )

                    if creada:
                        resumen[
                            "creadas"
                        ] += 1
                    else:
                        resumen[
                            "omitidas"
                        ] += 1

            except Exception as exc:
                resumen[
                    "errores"
                ] += 1

                self.stderr.write(
                    self.style.ERROR(
                        f"Fila {rownum}: {exc}"
                    )
                )


    # ========================================================
    # IMPORTAR FILA
    # ========================================================

    def importar_fila(
        self,
        workbook,
        sheet,
        path,
        row,
        rownum,
        headers,
        columnas,
        columnas_responsables,
        usuario,
        resumen,
    ):
        def get(campo):
            index = columnas.get(
                campo
            )

            if index is None:
                return None

            if index >= len(row):
                return None

            return row[index]


        # ----------------------------------------------------
        # Guardar TODA la fila original
        # ----------------------------------------------------

        originales = []

        for index, value in enumerate(
            row
        ):
            originales.append(
                {
                    "columna": index + 1,
                    "encabezado": (
                        headers[index]
                        if index < len(
                            headers
                        )
                        else ""
                    ),
                    "valor": serializable(
                        value
                    ),
                }
            )


        content_hash = hashlib.sha256(
            json.dumps(
                originales,
                ensure_ascii=False,
                sort_keys=True,
                default=str,
            ).encode(
                "utf-8"
            )
        ).hexdigest()


        identity = [
            FUENTE,
            sheet.title,
            rownum,
        ]


        clave = hashlib.sha256(
            json.dumps(
                identity,
                ensure_ascii=False,
            ).encode(
                "utf-8"
            )
        ).hexdigest()


        existente = (
            Tema.objects
            .filter(
                clave_importacion=clave
            )
            .first()
        )


        if existente:
            hash_anterior = (
                existente
                .datos_migracion
                .get(
                    "hash_contenido"
                )
            )

            if hash_anterior != content_hash:
                raise ValueError(
                    "La fila ya fue importada pero "
                    "el contenido del Excel cambió."
                )

            return False


        # ----------------------------------------------------
        # TIPO / DESCRIPCIÓN
        # ----------------------------------------------------

        tipo_original = texto(
            get("tipo")
        )

        informativo = texto(
            get("informativo")
        )

        problema = texto(
            get("problema")
        )


        tipo_normalizado = normalizar(
            tipo_original
        )


        if "problema" in tipo_normalizado:
            tipo_nombre = "Problema"

            descripcion = (
                problema
                or informativo
            )

        elif "inform" in tipo_normalizado:
            tipo_nombre = "Informativo"

            descripcion = (
                informativo
                or problema
            )

        else:
            # Fallback controlado.
            # Se usa el contenido de las columnas conocidas.

            if problema:
                tipo_nombre = "Problema"
                descripcion = problema

            elif informativo:
                tipo_nombre = "Informativo"
                descripcion = informativo

            else:
                raise ValueError(
                    "No puede determinarse el tipo "
                    "ni existe descripción."
                )


        tipo = self.resolver_tipo(
            tipo_nombre
        )


        if not descripcion:
            raise ValueError(
                "Descripción vacía."
            )


        titulo = generar_titulo(
            descripcion
        )


        impacto = texto(
            get("impacto")
        )


        # ----------------------------------------------------
        # ESTADO
        # ----------------------------------------------------

        estado_original = texto(
            get("estado")
        )

        estado = self.resolver_estado(
            estado_original
        )


        # ----------------------------------------------------
        # ÁREA
        # ----------------------------------------------------

        area_original = texto(
            get("area")
        )

        area = self.resolver_area(
            area_original,
            resumen,
        )


        # ----------------------------------------------------
        # RESPONSABLES MARCADOS
        # ----------------------------------------------------

        responsables_originales = []

        responsables_encontrados = []


        for nombre_historico, index in (
            columnas_responsables.items()
        ):
            if index >= len(row):
                continue

            if not celda_es_true(
                row[index]
            ):
                continue

            responsables_originales.append(
                nombre_historico
            )

            responsable = (
                self.resolver_responsable(
                    nombre_historico
                )
            )

            if responsable:
                responsables_encontrados.append(
                    responsable
                )
            else:
                resumen[
                    "responsables_no_encontrados"
                ].add(
                    nombre_historico
                )


        # Si hay exactamente uno podemos asignarlo
        # sin inventar una jerarquía.
        #
        # Si hay varios, quedan preservados históricamente
        # pero Tema.responsable queda sin asignar.

        responsable_principal = None

        responsables_unicos = {
            obj.pk: obj
            for obj
            in responsables_encontrados
        }

        if len(
            responsables_unicos
        ) == 1:
            responsable_principal = next(
                iter(
                    responsables_unicos.values()
                )
            )


        # ----------------------------------------------------
        # FECHAS
        # ----------------------------------------------------

        fecha_creacion = interpretar_fecha(
            get("fecha"),
            workbook.epoch,
        )


        fecha_limite, fecha_limite_texto = (
            intentar_fecha(
                get("fecha_limite"),
                workbook.epoch,
            )
        )


        # ----------------------------------------------------
        # REMITENTE
        # ----------------------------------------------------

        remitente_original = texto(
            get("remitente")
        )


        # ----------------------------------------------------
        # DATOS MIGRACIÓN
        # ----------------------------------------------------

        datos_migracion = {
            "fuente": FUENTE,
            "archivo": path.name,
            "hoja": sheet.title,
            "fila": rownum,

            "hash_contenido":
                content_hash,

            "tipo_original":
                tipo_original,

            "area_original":
                area_original,

            "estado_original":
                estado_original,

            "remitente_original":
                remitente_original,

            "responsables_historicos":
                responsables_originales,

            "responsables_encontrados": [
                responsable.nombre
                for responsable
                in responsables_encontrados
            ],

            "fecha_limite_original":
                texto(
                    get("fecha_limite")
                ),

            "celdas":
                originales,
        }


        # ----------------------------------------------------
        # CREAR TEMA
        # ----------------------------------------------------

        tema = Tema.objects.create(
            titulo=titulo,
            tipo=tipo,
            descripcion=descripcion,
            impacto=impacto,

            # El formulario histórico no tenía prioridad.
            prioridad=Tema.Prioridad.MEDIA,

            area_origen=area,

            responsable=responsable_principal,

            estado=estado,

            creador=usuario,

            clave_importacion=clave,

            datos_migracion=datos_migracion,
        )


        # Para las tarjetas del tablero resulta útil
        # que al menos el área de origen figure también
        # como área involucrada.

        tema.areas_involucradas.add(
            area
        )


        # ----------------------------------------------------
        # FECHAS HISTÓRICAS
        # ----------------------------------------------------

        updates = {}


        if fecha_creacion:
            updates[
                "fecha_creacion"
            ] = fecha_creacion

            updates[
                "fecha_actualizacion"
            ] = fecha_creacion


        if hasattr(
            Tema,
            "fecha_limite",
        ) and fecha_limite:
            updates[
                "fecha_limite"
            ] = fecha_limite


        if estado.es_resuelto:
            # El Excel no posee una fecha de resolución
            # confiable independiente.
            #
            # No inventamos una.
            updates[
                "fecha_resolucion"
            ] = None


        if updates:
            Tema.objects.filter(
                pk=tema.pk
            ).update(
                **updates
            )

            tema.refresh_from_db()


        # ----------------------------------------------------
        # ACTIVIDAD DE IMPORTACIÓN
        # ----------------------------------------------------

        ActividadTema.objects.create(
            tema=tema,
            usuario=usuario,
            tipo_evento="importacion",
            campo="",
            valor_anterior=None,
            valor_nuevo={
                "fuente": FUENTE,
                "hoja": sheet.title,
                "fila": rownum,
            },
        )


        # ----------------------------------------------------
        # COMENTARIOS HISTÓRICOS TIPIFICADOS
        # ----------------------------------------------------

        self.crear_comentario(
            tema=tema,
            usuario=usuario,
            texto_comentario=texto(
                get("solucion")
            ),
            tipo="solucion",
            resumen=resumen,
        )


        self.crear_comentario(
            tema=tema,
            usuario=usuario,
            texto_comentario=texto(
                get("estado_actual")
            ),
            tipo="seguimiento",
            resumen=resumen,
        )


        self.crear_comentario(
            tema=tema,
            usuario=usuario,
            texto_comentario=texto(
                get(
                    "comentario_reunion"
                )
            ),
            tipo="comentario",
            resumen=resumen,
        )


        # La fecha límite histórica a veces tiene texto
        # en vez de una fecha.
        #
        # Ese texto representa una acción/seguimiento,
        # por lo que se preserva visiblemente.

        if fecha_limite_texto:
            self.crear_comentario(
                tema=tema,
                usuario=usuario,
                texto_comentario=(
                    "Acción / fecha límite histórica:\n"
                    + fecha_limite_texto
                ),
                tipo="seguimiento",
                resumen=resumen,
            )


        return True


    # ========================================================
    # RESOLVER TIPO
    # ========================================================

    def resolver_tipo(
        self,
        nombre,
    ):
        candidatos = (
            TipoTema.objects
            .filter(
                activo=True
            )
        )


        buscado = normalizar(
            nombre
        )


        for tipo in candidatos:
            actual = normalizar(
                tipo.nombre
            )

            if actual == buscado:
                return tipo


        # Compatibilidad por si quedó cargado
        # "Tema Informativo" en alguna base.

        if buscado == "informativo":
            for tipo in candidatos:
                if "inform" in normalizar(
                    tipo.nombre
                ):
                    return tipo


        raise ValueError(
            f"No existe TipoTema compatible con {nombre!r}."
        )


    # ========================================================
    # RESOLVER ESTADO
    # ========================================================
    def resolver_estado(self, valor):
        """
        Traduce el estado histórico del Excel a los estados
        reales existentes en la base.

        No depende de que los códigos sean exactamente:
        nuevo / en-curso / resuelto.

        Usa:
        1. significado histórico;
        2. flags es_inicial / es_resuelto;
        3. nombre normalizado del maestro.
        """

        original = texto(valor)
        key = normalizar(original)

        estados = list(
            Estado.objects.filter(activo=True)
        )

        if not estados:
            raise ValueError(
                "No existen estados activos en la base de datos."
            )

        # ========================================================
        # DETERMINAR SIGNIFICADO DEL VALOR HISTÓRICO
        # ========================================================

        if not key or key in {
            "pendiente",
            "nuevo",
            "nuevos",
        }:
            significado = "nuevo"

        elif key in {
            "en curso",
            "curso",
            "en proceso",
        }:
            significado = "en_curso"

        elif key in {
            "resuelto",
            "resuelta",
            "resueltos",
            "finalizado",
            "finalizada",
            "finalizados",
        }:
            significado = "resuelto"

        else:
            raise ValueError(
                f"Estado histórico no reconocido: {original!r}"
            )

        # ========================================================
        # NUEVO
        # ========================================================

        if significado == "nuevo":

            # La definición estructural más confiable:
            # el estado inicial.
            for estado in estados:
                if estado.es_inicial:
                    return estado

            # Fallback por nombre/código.
            for estado in estados:
                nombre = normalizar(estado.nombre)
                codigo = normalizar(estado.codigo)

                if (
                    nombre in {"nuevo", "nuevos", "pendiente"}
                    or codigo in {"nuevo", "nuevos", "pendiente"}
                ):
                    return estado

            raise ValueError(
                "No se encontró un Estado activo que represente "
                "'Nuevo/Pendiente'."
            )

        # ========================================================
        # RESUELTO
        # ========================================================

        if significado == "resuelto":

            # El flag del modelo es la fuente más confiable.
            for estado in estados:
                if estado.es_resuelto:
                    return estado

            # Fallback por nombre/código.
            for estado in estados:
                nombre = normalizar(estado.nombre)
                codigo = normalizar(estado.codigo)

                if (
                    nombre in {
                        "resuelto",
                        "resueltos",
                        "finalizado",
                        "finalizados",
                    }
                    or codigo in {
                        "resuelto",
                        "resueltos",
                        "finalizado",
                        "finalizados",
                    }
                ):
                    return estado

            raise ValueError(
                "No se encontró un Estado activo que represente "
                "'Resuelto'."
            )

        # ========================================================
        # EN CURSO
        # ========================================================

        if significado == "en_curso":

            # Primero buscamos por semántica.
            for estado in estados:
                nombre = normalizar(estado.nombre)
                codigo = normalizar(estado.codigo)

                if (
                    nombre in {
                        "en curso",
                        "curso",
                        "en proceso",
                    }
                    or codigo in {
                        "en curso",
                        "en-curso",
                        "curso",
                        "en proceso",
                        "en-proceso",
                    }
                ):
                    return estado

            # Si exactamente tenemos:
            # - un inicial
            # - un resuelto
            # - un tercer estado
            #
            # ese tercero necesariamente es el estado intermedio.
            candidatos = [
                estado
                for estado in estados
                if not estado.es_inicial
                and not estado.es_resuelto
            ]

            if len(candidatos) == 1:
                return candidatos[0]

            raise ValueError(
                "No se pudo identificar de forma inequívoca "
                "el estado 'En curso'."
            )

        raise ValueError(
            f"No pudo resolverse el estado histórico {original!r}."
        )


    # ========================================================
    # RESOLVER ÁREA
    # ========================================================

    def resolver_area(
        self,
        valor,
        resumen,
    ):
        original = texto(
            valor
        )

        key = normalizar(
            original
        )


        if not key:
            raise ValueError(
                "Área de origen vacía."
            )


        # Primero buscamos alias de maestros existentes.

        destino = ALIASES_AREAS.get(
            key
        )


        if destino:
            area = self.buscar_area(
                destino
            )

            if area:
                return area


        # Luego intentamos match directo normalizado.

        for area in Area.objects.filter(
            activa=True
        ):
            if normalizar(
                area.nombre
            ) == key:
                return area


        # Hay áreas históricas realmente distintas
        # que no existían todavía como maestros.

        nueva = (
            AREAS_HISTORICAS_NUEVAS
            .get(
                key
            )
        )


        if nueva:
            area, created = (
                Area.objects.get_or_create(
                    nombre=nueva[
                        "nombre"
                    ],
                    defaults={
                        "color":
                            nueva[
                                "color"
                            ],
                        "activa":
                            True,
                    },
                )
            )

            if created:
                resumen[
                    "areas_creadas"
                ].append(
                    area.nombre
                )

            return area


        raise ValueError(
            f"Área histórica sin correspondencia: {original!r}"
        )


    def buscar_area(
        self,
        nombre,
    ):
        buscado = normalizar(
            nombre
        )

        for area in Area.objects.filter(
            activa=True
        ):
            if normalizar(
                area.nombre
            ) == buscado:
                return area

        return None


    # ========================================================
    # RESOLVER RESPONSABLE
    # ========================================================

    def resolver_responsable(
        self,
        valor,
    ):
        original = texto(
            valor
        )

        key = normalizar(
            original
        )


        # 1. Alias explícito

        destino = (
            ALIASES_RESPONSABLES
            .get(
                key
            )
        )


        if destino:
            buscado = normalizar(
                destino
            )

            for responsable in (
                Responsable.objects
                .filter(
                    activo=True
                )
            ):
                if normalizar(
                    responsable.nombre
                ) == buscado:
                    return responsable


        # 2. Match directo

        for responsable in (
            Responsable.objects
            .filter(
                activo=True
            )
        ):
            if normalizar(
                responsable.nombre
            ) == key:
                return responsable


        # 3. Invertir Apellido Nombre -> Nombre Apellido

        partes = original.split()

        if len(partes) >= 2:
            invertido = " ".join(
                partes[1:]
                + [partes[0]]
            )

            invertido_normalizado = (
                normalizar(
                    invertido
                )
            )

            for responsable in (
                Responsable.objects
                .filter(
                    activo=True
                )
            ):
                if normalizar(
                    responsable.nombre
                ) == invertido_normalizado:
                    return responsable


        # No se crean responsables automáticamente.
        return None


    # ========================================================
    # CREAR COMENTARIO HISTÓRICO
    # ========================================================

    def crear_comentario(
        self,
        tema,
        usuario,
        texto_comentario,
        tipo,
        resumen,
    ):
        texto_comentario = texto(
            texto_comentario
        )


        if not texto_comentario:
            return


        # Tu modelo admite 10.000 caracteres.
        #
        # Si existiera alguna respuesta más larga,
        # se conserva completa en fragmentos.

        for inicio in range(
            0,
            len(
                texto_comentario
            ),
            9500,
        ):
            fragmento = (
                texto_comentario[
                    inicio:
                    inicio + 9500
                ]
            )


            Comentario.objects.create(
                tema=tema,
                usuario=usuario,
                texto=fragmento,
                tipo=tipo,
                es_historico=True,
            )


            resumen[
                "comentarios"
            ] += 1