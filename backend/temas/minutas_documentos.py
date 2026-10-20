
"""
Generación de minutas institucionales en Word y PDF.

Fuente de datos:
    generar_datos_minuta(reunion)

Logo:
    temas/plantillas/logo_hospital.png

Los archivos se generan en memoria.
No se modifica la reunión ni su minuta histórica.
"""

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from django.core.exceptions import ImproperlyConfigured

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image,
    KeepTogether,
)

from .minutas import generar_datos_minuta


LOGO = Path(__file__).resolve().parent / "plantillas" / "logo_hospital.png"


def verificar_logo():
    if not LOGO.is_file():
        raise ImproperlyConfigured(
            f"No se encontró el logo institucional: {LOGO}"
        )


def datos_documento(reunion):
    """
    Una reunión confirmada usa exclusivamente
    la copia histórica almacenada.
    """

    if reunion.estado_minuta == "confirmada":
        if reunion.minuta_contenido is None:
            raise ValueError(
                "La reunión confirmada no tiene una minuta guardada."
            )
        return reunion.minuta_contenido

    return generar_datos_minuta(reunion)


def nombre_personas(personas, incluir_motivo=False):
    resultado = []

    for persona in personas:
        nombre = persona.get("nombre", "")

        if incluir_motivo:
            motivo = (persona.get("motivo") or "").strip()
            if motivo:
                nombre += f" ({motivo})"

        resultado.append(nombre)

    return ", ".join(resultado) if resultado else "Sin registrar"


def estado_tema(tema):
    estado = tema.get("estado_minuta")

    if not estado or estado == "Sin cambio de estado registrado":
        return "sin estado histórico determinado"

    return estado.lower()


def textos_tema(tema):
    """
    Conserva los comentarios registrados.
    No inventa decisiones ni compromisos.
    """

    textos = []

    for comentario in tema.get("comentarios", []):
        autor = comentario.get("autor", "Participante")
        texto = (comentario.get("texto") or "").strip()

        if texto:
            textos.append(f"{autor}: {texto}")

    if tema.get("cambio_estado"):
        anterior = (tema.get("estado_inicial") or {}).get("nombre")
        nuevo = (tema.get("estado_final") or {}).get("nombre")

        if anterior and nuevo:
            textos.append(
                f"Cambio de estado: {anterior} → {nuevo}."
            )

    return textos


def generar_word(reunion):
    verificar_logo()
    datos = datos_documento(reunion)

    documento = Document()
    seccion = documento.sections[0]

    seccion.page_width = Inches(8.5)
    seccion.page_height = Inches(11)
    seccion.top_margin = Inches(0.65)
    seccion.bottom_margin = Inches(0.65)
    seccion.left_margin = Inches(0.85)
    seccion.right_margin = Inches(0.85)

    normal = documento.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)

    encabezado = documento.add_paragraph()
    encabezado.alignment = WD_ALIGN_PARAGRAPH.LEFT
    encabezado.paragraph_format.space_after = Pt(5)
    encabezado.add_run().add_picture(
        str(LOGO),
        width=Inches(1.15),
    )

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(14)
    run = titulo.add_run("Minuta de reunión – Mandos Medios")
    run.bold = True
    run.font.size = Pt(12)

    p = documento.add_paragraph()
    p.add_run("Fecha: ").bold = True
    p.add_run(datos.get("fecha_formateada", ""))

    p = documento.add_paragraph()
    p.add_run("Asistentes: ").bold = True
    p.add_run(nombre_personas(datos.get("asistentes", [])))

    p = documento.add_paragraph()
    p.add_run("Ausentes: ").bold = True
    p.add_run(
        nombre_personas(
            datos.get("ausentes", []),
            incluir_motivo=True,
        )
    )

    sin_registrar = datos.get("sin_registrar", [])
    if sin_registrar:
        p = documento.add_paragraph()
        p.add_run("Asistencia sin registrar: ").bold = True
        p.add_run(nombre_personas(sin_registrar))

    p = documento.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run("Temas tratados:").bold = True

    for numero, tema in enumerate(datos.get("temas", []), start=1):
        p = documento.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.keep_with_next = True

        encabezado_tema = (
            f"{numero}. {tema.get('titulo', 'Sin título')} "
            f"({estado_tema(tema)})"
        )
        p.add_run(encabezado_tema).bold = True

        textos = textos_tema(tema)

        if not textos:
            documento.add_paragraph(
                "Sin comentarios registrados."
            )
        else:
            for texto in textos:
                p = documento.add_paragraph(texto)
                p.paragraph_format.left_indent = Inches(0.18)

    observaciones = (datos.get("observaciones") or "").strip()

    if observaciones:
        p = documento.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.add_run("Aclaración: ").bold = True
        p.add_run(observaciones)

    salida = BytesIO()
    documento.save(salida)
    salida.seek(0)
    return salida


def generar_pdf(reunion):
    verificar_logo()
    datos = datos_documento(reunion)

    salida = BytesIO()

    documento = SimpleDocTemplate(
        salida,
        pagesize=letter,
        rightMargin=61,
        leftMargin=61,
        topMargin=47,
        bottomMargin=47,
        title="Minuta de reunión – Mandos Medios",
    )

    estilo = ParagraphStyle(
        "NormalMinuta",
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        alignment=TA_LEFT,
        spaceAfter=8,
        textColor=colors.black,
    )

    titulo_estilo = ParagraphStyle(
        "TituloMinuta",
        parent=estilo,
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        alignment=TA_CENTER,
        spaceAfter=15,
    )

    tema_estilo = ParagraphStyle(
        "TemaMinuta",
        parent=estilo,
        fontName="Helvetica-Bold",
        spaceBefore=9,
        spaceAfter=6,
    )

    contenido_estilo = ParagraphStyle(
        "ContenidoMinuta",
        parent=estilo,
        leftIndent=13,
    )

    elementos = []

    imagen = ImageReader(str(LOGO))
    ancho, alto = imagen.getSize()
    ancho_final = 83
    alto_final = ancho_final * alto / ancho

    elementos.append(
        Image(
            str(LOGO),
            width=ancho_final,
            height=alto_final,
        )
    )

    elementos.append(Spacer(1, 6))

    elementos.append(
        Paragraph(
            "Minuta de reunión – Mandos Medios",
            titulo_estilo,
        )
    )

    elementos.append(
        Paragraph(
            f"<b>Fecha:</b> "
            f"{escape(datos.get('fecha_formateada', ''))}",
            estilo,
        )
    )

    elementos.append(
        Paragraph(
            "<b>Asistentes:</b> "
            + escape(nombre_personas(datos.get("asistentes", []))),
            estilo,
        )
    )

    elementos.append(
        Paragraph(
            "<b>Ausentes:</b> "
            + escape(
                nombre_personas(
                    datos.get("ausentes", []),
                    incluir_motivo=True,
                )
            ),
            estilo,
        )
    )

    sin_registrar = datos.get("sin_registrar", [])
    if sin_registrar:
        elementos.append(
            Paragraph(
                "<b>Asistencia sin registrar:</b> "
                + escape(nombre_personas(sin_registrar)),
                estilo,
            )
        )

    elementos.append(Spacer(1, 8))
    elementos.append(Paragraph("<b>Temas tratados:</b>", estilo))

    for numero, tema in enumerate(datos.get("temas", []), start=1):
        encabezado_tema = (
            f"{numero}. {tema.get('titulo', 'Sin título')} "
            f"({estado_tema(tema)})"
        )

        elementos_tema = [
            Paragraph(escape(encabezado_tema), tema_estilo)
        ]

        textos = textos_tema(tema)

        if not textos:
            elementos_tema.append(
                Paragraph(
                    "Sin comentarios registrados.",
                    contenido_estilo,
                )
            )
        else:
            for texto in textos:
                elementos_tema.append(
                    Paragraph(
                        escape(texto).replace("\n", "<br/>"),
                        contenido_estilo,
                    )
                )

        # Mantener el título junto al primer párrafo,
        # pero permitir que los temas largos ocupen
        # más de una página.
        elementos.append(
            KeepTogether(elementos_tema[:2])
        )
        elementos.extend(elementos_tema[2:])

    observaciones = (datos.get("observaciones") or "").strip()

    if observaciones:
        elementos.append(Spacer(1, 12))
        elementos.append(
            Paragraph(
                "<b>Aclaración:</b> "
                + escape(observaciones).replace("\n", "<br/>"),
                estilo,
            )
        )

    documento.build(elementos)
    salida.seek(0)
    return salida
