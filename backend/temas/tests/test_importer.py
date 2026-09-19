import json
from datetime import datetime
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from temas.models import Area, Comentario, Responsable, Tema


class ImporterTest(TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.file = self.root / "historico.xlsx"
        self.mapping = self.root / "mapping.json"
        self.config = {
            "columnas": {
                "id_origen": "ID",
                "descripcion": "Planteo",
                "area_origen": "Sector",
                "fecha_creacion": "Fecha",
                "estado": "Estado",
                "tipo": "Tipo",
                "responsable": "Responsable",
            },
            "comentarios": ["Seguimiento"],
            "predeterminados": {"prioridad": "Media"},
            "valores": {
                "estado": {"Pendiente": "nuevo", "Finalizado": "resuelto"},
                "tipo": {"Problema": "Problema"},
                "prioridad": {"Media": "media"},
            },
        }
        self.mapping.write_text(json.dumps(self.config))
        get_user_model().objects.create_user("importador")
        call_command("cargar_iniciales", stdout=StringIO())
        self.headers = [
            "ID",
            "Planteo",
            "Sector",
            "Fecha",
            "Estado",
            "Tipo",
            "Responsable",
            "Seguimiento",
            "Columna no mapeada",
        ]
        self.row = [
            "A1",
            "Demora en camas",
            "Enfermería",
            datetime(2024, 5, 10, 12, 30),
            "Pendiente",
            "Problema",
            "Laura",
            "Se revisó en reunión.",
            "Dato que debe conservarse",
        ]
        self.book([self.row])

    def book(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.title = "Histórico"
        sheet.append(self.headers)
        for row in rows:
            sheet.append(row)
        book.save(self.file)

    def run_import(self, **extra):
        call_command(
            "importar_temas",
            str(self.file),
            mapping=str(self.mapping),
            usuario="importador",
            fuente="prueba-historica",
            stdout=StringIO(),
            stderr=StringIO(),
            **extra,
        )

    def test_all_cells_preserved_and_historical_comments(self):
        self.run_import()
        tema = Tema.objects.get()
        self.assertEqual(tema.fecha_creacion.year, 2024)
        self.assertEqual(tema.datos_migracion["celdas"][-1]["valor"], self.row[-1])
        self.assertEqual(tema.creador.username, "importador")
        self.assertTrue(Comentario.objects.get().es_historico)

    def test_dry_run_rolls_back_topics_masters_and_comments(self):
        self.run_import(dry_run=True)
        for model in [Tema, Area, Responsable, Comentario]:
            self.assertEqual(model.objects.count(), 0)

    def test_idempotence(self):
        self.run_import()
        self.run_import()
        self.assertEqual(Tema.objects.count(), 1)
        self.assertEqual(Comentario.objects.count(), 1)

    def test_existing_modified_row_rejected_without_overwrite(self):
        self.run_import()
        self.row[1] = "Texto diferente"
        self.book([self.row])
        with self.assertRaises(CommandError):
            self.run_import()
        self.assertEqual(Tema.objects.get().descripcion, "Demora en camas")

    def test_invalid_row_does_not_prevent_valid_and_rolls_back_masters(self):
        bad = list(self.row)
        bad[0] = "A2"
        bad[2] = "Área errónea"
        bad[3] = "fecha imposible"
        self.book([bad, self.row])
        with self.assertRaises(CommandError):
            self.run_import()
        self.assertEqual(Tema.objects.count(), 1)
        self.assertFalse(Area.objects.filter(nombre="Área errónea").exists())

    def test_unknown_state_rejected(self):
        self.row[4] = "???"
        self.book([self.row])
        with self.assertRaises(CommandError):
            self.run_import()
        self.assertEqual(Tema.objects.count(), 0)

    def test_unknown_resolution_date_not_invented(self):
        self.row[4] = "Finalizado"
        self.book([self.row])
        self.run_import()
        self.assertIsNone(Tema.objects.get().fecha_resolucion)

    def test_missing_header_and_duplicate_header_rejected(self):
        self.headers[1] = "Otro"
        self.book([self.row])
        with self.assertRaises(CommandError):
            self.run_import()
        self.headers[1] = "Sector"
        self.book([self.row])
        with self.assertRaises(CommandError):
            self.run_import()

    def test_formula_rejected_preserving_no_partial_topic(self):
        self.row[1] = "=A1"
        self.book([self.row])
        with self.assertRaises(CommandError):
            self.run_import()
        self.assertEqual(Tema.objects.count(), 0)

    def test_inspection_visits_all_sheets_columns_rows(self):
        book = Workbook()
        for sheet in [book.active, book.create_sheet("Otra")]:
            sheet.append(["A", "B"])
            sheet.append([1, None])
            sheet.append(["texto", 5])
        book.save(self.file)
        report = self.root / "inspect.json"
        call_command(
            "importar_temas", str(self.file), inspect=True, report=str(report), stdout=StringIO()
        )
        data = json.loads(report.read_text())
        self.assertEqual(len(data["hojas"]), 2)
        self.assertEqual(data["hojas"][1]["datos"], [[1, None], ["texto", 5]])
        self.assertEqual(Tema.objects.count(), 0)

    def test_missing_file_and_missing_mapping_rejected(self):
        with self.assertRaises(CommandError):
            call_command("importar_temas", "no-existe.xlsx")
        with self.assertRaises(CommandError):
            call_command("importar_temas", str(self.file))
